"""Excel エクスポートモジュール"""
from typing import List, Dict
from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


class ExcelExporter:
    """名刺データをExcel形式でエクスポートするクラス"""

    @staticmethod
    def create_excel_from_dataframe(df: pd.DataFrame, sheet_name: str = "名刺データ") -> BytesIO:
        """
        DataFrameからExcelファイルを作成

        Args:
            df: エクスポートするDataFrame
            sheet_name: シート名

        Returns:
            Excelファイルのバイナリデータ
        """
        # 列名を日本語に変換
        column_names = {
            "name": "氏名",
            "name_kana": "氏名（かな）",
            "company": "会社名",
            "department": "部署",
            "position": "役職",
            "postal_code": "郵便番号",
            "address": "住所",
            "phone": "電話番号",
            "mobile": "携帯電話",
            "fax": "FAX",
            "email": "メールアドレス",
            "website": "ウェブサイト"
        }

        # DataFrameの列名を変換
        df_export = df.copy()
        df_export = df_export.rename(columns=column_names)

        # 内部フィールド（_で始まる列）を除外
        df_export = df_export[[col for col in df_export.columns if not col.startswith('_')]]

        # BytesIOオブジェクトを作成
        output = BytesIO()

        # pandasを使ってExcelに書き込み（文字化け防止）
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df_export.to_excel(writer, sheet_name=sheet_name, index=False)

        # スタイルを適用するためにワークブックを再読み込み
        output.seek(0)
        wb = load_workbook(output)
        ws = wb[sheet_name]

        # ヘッダー行のスタイル設定
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF", name="メイリオ")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # すべてのセルに罫線とスタイルを追加
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = thin_border
                # データ行のスタイル設定
                if cell.row > 1:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    # 日本語フォントを指定
                    cell.font = Font(name="メイリオ")

        # 列幅を自動調整
        for column_cells in ws.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter

            for cell in column_cells:
                try:
                    if cell.value:
                        # 日本語文字を考慮した幅計算
                        cell_value = str(cell.value)
                        # 日本語は1文字=2、英数字は1文字=1として計算
                        length = sum(2 if ord(c) > 127 else 1 for c in cell_value)
                        if length > max_length:
                            max_length = length
                except:
                    pass

            # 最小幅10、最大幅50
            adjusted_width = min(max(max_length + 2, 10), 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # 先頭行を固定
        ws.freeze_panes = "A2"

        # BytesIOに保存
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return output

    @staticmethod
    def create_excel_from_cards(cards: List[Dict[str, str]], sheet_name: str = "名刺データ") -> BytesIO:
        """
        名刺データのリストからExcelファイルを作成

        Args:
            cards: 名刺データのリスト
            sheet_name: シート名

        Returns:
            Excelファイルのバイナリデータ
        """
        if not cards:
            # 空のDataFrameを作成
            df = pd.DataFrame(columns=[
                "name", "name_kana", "company", "department", "position",
                "postal_code", "address", "phone", "mobile", "fax", "email", "website"
            ])
        else:
            df = pd.DataFrame(cards)

        return ExcelExporter.create_excel_from_dataframe(df, sheet_name)
